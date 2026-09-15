#!/usr/bin/env python3
"""Render the client User Guide: dist/USER_GUIDE.md -> USER_GUIDE.html + USER_GUIDE.docx.

    <python with python-docx + Pillow> tools/build_guide.py               # -> dist/USER_GUIDE.html/.docx
    <python with python-docx + Pillow> tools/build_guide.py --out <dir>   # review copy elsewhere
    python3 tools/build_guide.py --html-only --out <dir>                   # no python-docx needed

dist/USER_GUIDE.md and the screenshots it references (dist/images/) are the source of truth.
This script only READS them: it never creates, deletes or overwrites anything in dist/images/
and never rewrites the .md. It writes exactly USER_GUIDE.html and USER_GUIDE.docx into --out.

HTML: one self-contained file (inline CSS, images as base64, no external requests), heading ids
and a table of contents before the first "## " section. DOCX: python-docx, images embedded,
headings bookmarked and linked from a contents list.
Exits 1 and writes nothing if the guide references an image that does not exist.

Setup: python3 -m venv <venv> && <venv>/bin/pip install python-docx Pillow
"""
from __future__ import annotations

import argparse
import base64
import html
import itertools
import mimetypes
import os
import re
import sys
from collections import namedtuple
from urllib.parse import unquote

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")

# =================================================================== parsing ==
HEAD_RE = re.compile(r"^ {0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
IMG_RE = re.compile(r'^\s*!\[(?P<alt>[^\]]*)\]\((?P<src>[^)\s]+)(?:\s+"[^"]*")?\)\s*$')
LIST_RE = re.compile(r"^(?P<ind> *)(?P<mark>[-*+]|\d{1,9}[.)]) +(?P<text>\S.*)$")
FENCE_RE = re.compile(r"^\s*(```+|~~~+)\s*([\w+-]*)\s*$")
HR_RE = re.compile(r"^ {0,3}([-*_])( *\1){2,} *$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")


def indent_of(line):
    return len(line) - len(line.lstrip(" "))


def starts_block(line):
    s = line.lstrip()
    return bool(HR_RE.match(line) or HEAD_RE.match(line) or FENCE_RE.match(line)
                or IMG_RE.match(line) or s.startswith(">") or s.startswith("|")
                or LIST_RE.match(line))


def is_ordered(mark):
    return mark[0].isdigit()


def split_row(row):
    s = row.strip()
    s = s[1:] if s.startswith("|") else s
    s = s[:-1] if s.endswith("|") and not s.endswith("\\|") else s
    cells, cur, in_code, k = [], "", False, 0
    while k < len(s):
        ch = s[k]
        if ch == "\\" and s[k + 1:k + 2] == "|":
            cur += "|"; k += 2; continue
        if ch == "`":
            in_code = not in_code
        if ch == "|" and not in_code:
            cells.append(cur.strip()); cur = ""
        else:
            cur += ch
        k += 1
    cells.append(cur.strip())
    return cells


def table_block(rows):
    aligns = []
    for c in split_row(rows[1]):
        aligns.append("center" if c.startswith(":") and c.endswith(":")
                      else "right" if c.endswith(":") else "left")
    return ("table", split_row(rows[0]), aligns, [split_row(r) for r in rows[2:]])


def parse_list(lines, i):
    """-> (("list", ordered, start, [(text, [child blocks])]), next_index)"""
    n = len(lines)
    first = LIST_RE.match(lines[i])
    base = len(first.group("ind"))
    ordered = is_ordered(first.group("mark"))
    start = int(first.group("mark")[:-1]) if ordered else 1
    items = []
    while i < n:
        m = LIST_RE.match(lines[i])
        if not m or len(m.group("ind")) != base or is_ordered(m.group("mark")) != ordered:
            break
        text, children, after_blank = [m.group("text").strip()], [], False
        i += 1
        while i < n:
            line = lines[i]
            if not line.strip():
                j = i
                while j < n and not lines[j].strip():
                    j += 1
                if j < n and indent_of(lines[j]) > base:
                    i, after_blank = j, True
                    continue
                break
            lm = LIST_RE.match(line)
            if indent_of(line) <= base:
                if lm or after_blank or starts_block(line):
                    break
                text.append(line.strip()); i += 1           # lazy continuation
                continue
            if lm:
                sub, i = parse_list(lines, i)
                children.append(sub); after_blank = False
                continue
            im = IMG_RE.match(line)
            if im:
                children.append(("img", im.group("alt"), im.group("src")))
                i += 1; after_blank = False
                continue
            if after_blank or children:
                buf = []
                while (i < n and lines[i].strip() and indent_of(lines[i]) > base
                       and not LIST_RE.match(lines[i]) and not IMG_RE.match(lines[i])):
                    buf.append(lines[i].strip()); i += 1
                children.append(("p", " ".join(buf))); after_blank = False
                continue
            text.append(line.strip()); i += 1
        items.append((" ".join(text), children))
        j = i
        while j < n and not lines[j].strip():
            j += 1
        m2 = LIST_RE.match(lines[j]) if j < n else None
        if m2 and len(m2.group("ind")) == base and is_ordered(m2.group("mark")) == ordered:
            i = j
            continue
        break
    return ("list", ordered, start, items), i


def parse_blocks(lines):
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1; continue
        m = FENCE_RE.match(line)
        if m:
            fence, buf = m.group(1)[:3], []
            i += 1
            while i < n and not lines[i].strip().startswith(fence):
                buf.append(lines[i]); i += 1
            blocks.append(("code", "\n".join(buf))); i += 1
            continue
        if HR_RE.match(line):
            blocks.append(("hr",)); i += 1; continue
        m = HEAD_RE.match(line)
        if m:
            blocks.append(("heading", len(m.group(1)), m.group(2))); i += 1; continue
        m = IMG_RE.match(line)
        if m:
            blocks.append(("img", m.group("alt"), m.group("src"))); i += 1; continue
        if line.lstrip().startswith(">"):
            paras, cur = [], []
            while i < n and lines[i].lstrip().startswith(">"):
                t = re.sub(r"^\s*> ?", "", lines[i]).strip()
                if t:
                    cur.append(t)
                elif cur:
                    paras.append(" ".join(cur)); cur = []
                i += 1
            if cur:
                paras.append(" ".join(cur))
            blocks.append(("quote", paras))
            continue
        if line.lstrip().startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(lines[i]); i += 1
            blocks.append(table_block(rows))
            continue
        if LIST_RE.match(line):
            block, i = parse_list(lines, i)
            blocks.append(block)
            continue
        buf = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not starts_block(lines[i]):
            buf.append(lines[i].strip()); i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


# ------------------------------------------------------------------- inline --
Run = namedtuple("Run", "text bold italic code href")
CODE_RE = re.compile(r"(`+)(.+?)\1")
TOKEN_RE = re.compile(
    r"(?P<code>\x00(?P<ci>\d+)\x00)"
    r"|(?P<link>\[(?P<ltext>[^\]]+)\]\((?P<href>[^)\s]+)\))"
    r"|(?P<bold>\*\*(?=\S)(?P<btext>.+?)(?<=\S)\*\*)"
    r"|(?P<ital>(?<![\w*])\*(?=[^\s*])(?P<itext>.+?)(?<=[^\s*])\*(?![\w*]))")


def inline_runs(text):
    """Markdown inline -> [Run]. Code spans are protected first, so `code` works inside
    **bold**, *italic* and [links], and * inside code is literal."""
    codes = []

    def stash(m):
        codes.append(m.group(2).strip() or m.group(2))
        return "\x00%d\x00" % (len(codes) - 1)

    def runs(t, bold, italic, href):
        out, pos = [], 0
        for m in TOKEN_RE.finditer(t):
            if m.start() > pos:
                out.append(Run(t[pos:m.start()], bold, italic, False, href))
            if m.group("code"):
                out.append(Run(codes[int(m.group("ci"))], bold, italic, True, href))
            elif m.group("link"):
                out += runs(m.group("ltext"), bold, italic, m.group("href"))
            elif m.group("bold"):
                out += runs(m.group("btext"), True, italic, href)
            else:
                out += runs(m.group("itext"), bold, True, href)
            pos = m.end()
        if pos < len(t):
            out.append(Run(t[pos:], bold, italic, False, href))
        return out

    return runs(CODE_RE.sub(stash, text), False, False, None)


def plain(text):
    return "".join(r.text for r in inline_runs(text))


def slugify(text, used):
    s = re.sub(r"[^\w\s-]", "", plain(text).lower()).strip()
    s = re.sub(r"[\s_-]+", "-", s).strip("-") or "section"
    base, k = s, 2
    while s in used:
        s, k = "%s-%d" % (base, k), k + 1
    used.add(s)
    return s


def assign_ids(blocks):
    used = set()
    return [("heading", b[1], b[2], slugify(b[2], used)) if b[0] == "heading" else b
            for b in blocks]


def iter_images(blocks):
    for b in blocks:
        if b[0] == "img":
            yield b
        elif b[0] == "list":
            for _text, children in b[3]:
                yield from iter_images(children)


def resolve(base_dir, src):
    return os.path.normpath(os.path.join(base_dir, unquote(src)))


# ====================================================================== HTML ==
CSS = """
:root{--fg:#1d2330;--muted:#5b6475;--bg:#fff;--line:#dde2ea;--accent:#2e6fd1;--quote:#f1f6fd;
--code:#eef1f5;--th:#2e6fd1;--thfg:#fff;--zebra:#f7f9fb}
@media (prefers-color-scheme:dark){:root{--fg:#e6e9ef;--muted:#a3abba;--bg:#14171c;--line:#2c323c;
--accent:#79a8ff;--quote:#1b2433;--code:#232833;--th:#24406e;--thfg:#fff;--zebra:#191d23}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.6 -apple-system,"Segoe UI",
Helvetica,Arial,sans-serif;-webkit-text-size-adjust:100%}
.guide{max-width:880px;margin:0 auto;padding:32px 20px 64px}
h1{font-size:2rem;line-height:1.25;margin:0 0 16px;padding-bottom:10px;border-bottom:3px solid var(--accent)}
h2{font-size:1.45rem;line-height:1.3;margin:40px 0 12px;padding-bottom:6px;border-bottom:1px solid var(--line)}
h3{font-size:1.15rem;margin:28px 0 8px}
h1,h2,h3,h4{scroll-margin-top:16px}
a{color:var(--accent)}
img{max-width:100%;height:auto;display:block;margin:0 auto;border:1px solid var(--line);border-radius:6px}
figure{margin:16px 0}
figcaption{text-align:center;color:var(--muted);font-size:.8rem;font-style:italic;margin-top:6px}
blockquote{margin:16px 0;padding:8px 16px;border-left:4px solid var(--accent);background:var(--quote);
border-radius:0 6px 6px 0}
blockquote p{margin:6px 0}
code{font-family:SFMono-Regular,Consolas,Menlo,monospace;font-size:.88em;background:var(--code);
padding:1px 5px;border-radius:4px}
pre{background:var(--code);padding:12px 14px;border-radius:6px;overflow-x:auto}
pre code{background:none;padding:0}
.table-wrap{overflow-x:auto;margin:16px 0}
table{border-collapse:collapse;width:100%;font-size:.92rem}
th,td{border:1px solid var(--line);padding:7px 10px;text-align:left;vertical-align:top}
th{background:var(--th);color:var(--thfg)}
tbody tr:nth-child(even){background:var(--zebra)}
hr{border:none;border-top:1px solid var(--line);margin:32px 0}
li{margin:4px 0}
li>figure{margin:10px 0}
.toc{margin:24px 0;padding:14px 20px;border:1px solid var(--line);border-radius:8px}
.toc-title{margin:0 0 6px;font-weight:600}
.toc ul{margin:0;padding-left:20px}
.toc li{margin:2px 0}
@media print{.toc a{color:inherit;text-decoration:none}figure{break-inside:avoid}}
"""


def runs_html(runs):
    parts = []
    for (href, bold, italic), group in itertools.groupby(runs, lambda r: (r.href, r.bold, r.italic)):
        inner = "".join("<code>%s</code>" % html.escape(r.text, quote=False) if r.code
                        else html.escape(r.text, quote=False) for r in group)
        if italic:
            inner = "<em>%s</em>" % inner
        if bold:
            inner = "<strong>%s</strong>" % inner
        if href:
            inner = '<a href="%s">%s</a>' % (html.escape(href), inner)
        parts.append(inner)
    return "".join(parts)


class HtmlWriter:
    def __init__(self, base_dir):
        self.base_dir = base_dir
        self.cache = {}

    def data_uri(self, src):
        path = resolve(self.base_dir, src)
        if path not in self.cache:
            mime = mimetypes.guess_type(path)[0] or "application/octet-stream"
            with open(path, "rb") as f:
                self.cache[path] = "data:%s;base64,%s" % (mime, base64.b64encode(f.read()).decode("ascii"))
        return self.cache[path]

    def toc(self, blocks):
        groups = []
        for b in blocks:
            if b[0] != "heading" or b[1] not in (2, 3):
                continue
            if b[1] == 2 or not groups:
                groups.append([b if b[1] == 2 else None, [] if b[1] == 2 else [b]])
            else:
                groups[-1][1].append(b)
        if not groups:
            return ""

        def link(b):
            return '<a href="#%s">%s</a>' % (b[3], html.escape(plain(b[2]).strip(), quote=False))

        out = ['<nav class="toc" aria-label="Contents"><p class="toc-title">Contents</p><ul>']
        for head, subs in groups:
            sub = "<ul>%s</ul>" % "".join("<li>%s</li>" % link(s) for s in subs) if subs else ""
            out.append("<li>%s%s</li>" % (link(head) if head else "", sub))
        out.append("</ul></nav>")
        return "".join(out)

    def block(self, b):
        kind = b[0]
        if kind == "heading":
            lvl = min(b[1], 6)
            return '<h%d id="%s">%s</h%d>' % (lvl, b[3], runs_html(inline_runs(b[2])), lvl)
        if kind == "p":
            return "<p>%s</p>" % runs_html(inline_runs(b[1]))
        if kind == "quote":
            return "<blockquote>%s</blockquote>" % "".join(
                "<p>%s</p>" % runs_html(inline_runs(p)) for p in b[1])
        if kind == "hr":
            return "<hr>"
        if kind == "code":
            return "<pre><code>%s</code></pre>" % html.escape(b[1], quote=False)
        if kind == "img":
            alt = html.escape(b[1])
            return ('<figure><img alt="%s" src="%s"><figcaption>%s</figcaption></figure>'
                    % (alt, self.data_uri(b[2]), alt))
        if kind == "table":
            _, header, aligns, rows = b

            def cell(tag, j, text):
                al = aligns[j] if j < len(aligns) else "left"
                style = "" if al == "left" else ' style="text-align:%s"' % al
                return "<%s%s>%s</%s>" % (tag, style, runs_html(inline_runs(text)), tag)

            thead = "".join(cell("th", j, c) for j, c in enumerate(header))
            tbody = "".join("<tr>%s</tr>" % "".join(cell("td", j, c) for j, c in
                                                    enumerate(r[:len(header)]))
                            for r in rows)
            return ('<div class="table-wrap"><table><thead><tr>%s</tr></thead>'
                    '<tbody>%s</tbody></table></div>' % (thead, tbody))
        if kind == "list":
            _, ordered, start, items = b
            tag = "ol" if ordered else "ul"
            attr = ' start="%d"' % start if ordered and start != 1 else ""
            lis = "".join("<li>%s%s</li>" % (runs_html(inline_runs(text)),
                                            "".join(self.block(c) for c in children))
                          for text, children in items)
            return "<%s%s>%s</%s>" % (tag, attr, lis, tag)
        raise ValueError(kind)

    def document(self, blocks, title):
        body, toc_done, toc = [], False, self.toc(blocks)
        for b in blocks:
            if not toc_done and toc and b[0] == "heading" and b[1] == 2:
                body.append(toc); toc_done = True
            body.append(self.block(b))
        return ('<!doctype html>\n<html lang="en"><head><meta charset="utf-8">\n'
                '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
                "<title>%s</title>\n<style>%s</style>\n</head>\n<body><main class=\"guide\">\n%s\n"
                "</main></body></html>\n" % (html.escape(title), CSS, "\n".join(body)))


# ====================================================================== DOCX ==
def render_docx(blocks, base_dir, out_path, title):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from PIL import Image

    ACCENT, GREY = "2E6FD1", RGBColor(0x6B, 0x72, 0x80)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    usable_in = (sec.page_width - sec.left_margin - sec.right_margin) / 914400.0
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    cp = doc.core_properties
    cp.title, cp.author, cp.last_modified_by = title, "BYTEPLUS for Maya", "BYTEPLUS for Maya"
    cp.comments = cp.subject = cp.keywords = ""

    heads = [b for b in blocks if b[0] == "heading" and b[1] in (2, 3)]
    anchors = {b[3]: "_Toc_bp%03d" % k for k, b in enumerate(heads, 1)}
    bm_ids = itertools.count(1)

    def xml(tag, **attrs):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn("w:" + k), str(v))
        return el

    def shade(parent, fill):
        parent.append(xml("w:shd", val="clear", color="auto", fill=fill))

    def add_link(p, text, anchor=None, url=None):
        h = OxmlElement("w:hyperlink")
        if url:
            h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        else:
            h.set(qn("w:anchor"), anchor)
        h.set(qn("w:history"), "1")
        r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        rpr.append(xml("w:color", val=ACCENT))
        rpr.append(xml("w:u", val="single"))
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(rpr); r.append(t); h.append(r)
        p._p.append(h)

    def add_runs(p, runs, size=None, force_bold=False):
        for r in runs:
            if r.href:
                add_link(p, r.text, url=r.href)
                continue
            run = p.add_run(r.text)
            if r.bold or force_bold:
                run.bold = True
            if r.italic:
                run.italic = True
            if r.code:
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2B, 0x2F, 0x36)
                shade(run._r.get_or_add_rPr(), "EEF1F5")
            elif size:
                run.font.size = size

    def bordered(p, side, fill=None):
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bdr.append(xml("w:" + side, val="single", sz=18 if side == "left" else 6,
                       space=8 if side == "left" else 1,
                       color=ACCENT if side == "left" else "D0D5DD"))
        ppr.append(bdr)
        if fill:
            shade(ppr, fill)

    def picture(src, alt, indent_in=0.0):
        path = resolve(base_dir, src)
        with Image.open(path) as im:
            wpx, hpx = im.size
        w = min(usable_in - indent_in, wpx / 96.0)
        if hpx * w / wpx > 8.0:
            w = 8.0 * wpx / hpx
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if indent_in:
            p.paragraph_format.left_indent = Inches(indent_in)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(w))
        cap = doc.add_paragraph()
        if indent_in:
            cap.paragraph_format.left_indent = Inches(indent_in)
        else:
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(alt)
        cr.italic = True
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = GREY

    def add_list(block, level):
        _, ordered, start, items = block
        left = 0.3 + 0.3 * level
        for k, (text, children) in enumerate(items):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(left), Inches(-0.25)
            pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Inches(left))
            p.add_run(("%d." % (start + k)) if ordered else ("•" if level % 2 == 0 else "◦"))
            p.add_run("\t")
            add_runs(p, inline_runs(text))
            for c in children:
                if c[0] == "list":
                    add_list(c, level + 1)
                elif c[0] == "img":
                    picture(c[2], c[1], indent_in=left)
                elif c[0] == "p":
                    q = doc.add_paragraph()
                    q.paragraph_format.left_indent = Inches(left)
                    add_runs(q, inline_runs(c[1]))
                else:
                    render(c)

    def toc():
        p = doc.add_paragraph()
        r = p.add_run("Contents")
        r.bold = True
        r.font.size = Pt(13)
        for b in heads:
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            q.paragraph_format.left_indent = Inches(0 if b[1] == 2 else 0.3)
            add_link(q, plain(b[2]).strip(), anchor=anchors[b[3]])

    def render(b):
        kind = b[0]
        if kind == "heading":
            p = doc.add_heading("", {1: 0, 2: 1, 3: 2}.get(b[1], 3))
            add_runs(p, inline_runs(b[2]))
            if b[3] in anchors:
                i = str(next(bm_ids))
                start = xml("w:bookmarkStart", id=i, name=anchors[b[3]])
                ppr = p._p.pPr
                if ppr is not None:
                    ppr.addnext(start)
                else:
                    p._p.insert(0, start)
                p._p.append(xml("w:bookmarkEnd", id=i))
        elif kind == "p":
            add_runs(doc.add_paragraph(), inline_runs(b[1]))
        elif kind == "quote":
            for para in b[1]:
                p = doc.add_paragraph()
                bordered(p, "left", fill="F1F6FD")
                p.paragraph_format.left_indent = Inches(0.12)
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
                add_runs(p, inline_runs(para))
        elif kind == "hr":
            bordered(doc.add_paragraph(), "bottom")
        elif kind == "code":
            for ln in b[1].split("\n") or [""]:
                p = doc.add_paragraph()
                shade(p._p.get_or_add_pPr(), "EEF1F5")
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(ln)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif kind == "img":
            picture(b[2], b[1])
        elif kind == "table":
            _, header, _aligns, rows = b
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, c in enumerate(header):
                add_runs(t.rows[0].cells[j].paragraphs[0], inline_runs(c), force_bold=True)
            for row in rows:
                cells = t.add_row().cells
                for j, c in enumerate(row[:len(header)]):
                    add_runs(cells[j].paragraphs[0], inline_runs(c))
        elif kind == "list":
            add_list(b, 0)

    toc_done = False
    for b in blocks:
        if not toc_done and heads and b[0] == "heading" and b[1] == 2:
            toc(); toc_done = True
        render(b)
    tmp = out_path + ".part"
    doc.save(tmp)
    os.replace(tmp, out_path)


# ====================================================================== main ==
def plugin_version():
    try:
        with open(os.path.join(ROOT, "byteplus_maya.py"), encoding="utf-8") as f:
            m = re.search(r'^\s{4}VERSION\s*=\s*"([^"]+)"', f.read(), re.M)
        return m.group(1).split()[0] if m else None
    except OSError:
        return None


def main(argv=None):
    ap = argparse.ArgumentParser(description="Render USER_GUIDE.md to self-contained HTML + DOCX "
                                             "(never touches the images folder).")
    ap.add_argument("--md", default=os.path.join(DIST, "USER_GUIDE.md"), help="guide source")
    ap.add_argument("--out", default=DIST, help="folder for USER_GUIDE.html/.docx (default: dist/)")
    ap.add_argument("--html-only", action="store_true", help="skip the DOCX (no python-docx needed)")
    args = ap.parse_args(argv)

    md = os.path.abspath(args.md)
    base_dir = os.path.dirname(md)
    out = os.path.abspath(args.out)
    images_dir = os.path.realpath(os.path.join(base_dir, "images"))
    real_out = os.path.realpath(out)
    if real_out == images_dir or real_out.startswith(images_dir + os.sep):
        print("refusing to write into the images folder: %s" % out)
        return 2

    with open(md, encoding="utf-8") as f:
        lines = [ln.expandtabs(4) for ln in f.read().splitlines()]
    blocks = assign_ids(parse_blocks(lines))
    title = next((plain(b[2]).strip() for b in blocks if b[0] == "heading" and b[1] == 1),
                 "User Guide")

    imgs = list(iter_images(blocks))
    bad = [src for _k, _alt, src in imgs
           if re.match(r"^[a-z][a-z0-9+.-]*:", src, re.I) or not os.path.isfile(resolve(base_dir, src))]
    if bad:
        print("missing or non-local images referenced by %s (nothing written):" % md)
        for src in bad:
            print("  - " + src)
        return 1
    if not args.html_only:
        try:
            import docx  # noqa: F401
            import PIL  # noqa: F401
        except ImportError as e:
            print("python-docx and Pillow are required for the DOCX (%s); "
                  "use a venv with `pip install python-docx Pillow`, or --html-only." % e)
            return 2

    os.makedirs(out, exist_ok=True)
    html_path = os.path.join(out, "USER_GUIDE.html")
    with open(html_path + ".part", "w", encoding="utf-8", newline="\n") as f:
        f.write(HtmlWriter(base_dir).document(blocks, title))
    os.replace(html_path + ".part", html_path)
    print("wrote %s (%d bytes)" % (html_path, os.path.getsize(html_path)))
    if not args.html_only:
        docx_path = os.path.join(out, "USER_GUIDE.docx")
        render_docx(blocks, base_dir, docx_path, title)
        print("wrote %s (%d bytes)" % (docx_path, os.path.getsize(docx_path)))

    heads = sum(1 for b in blocks if b[0] == "heading")
    print("%d headings, %d image references (%d files), images read from %s (unchanged)"
          % (heads, len(imgs), len({resolve(base_dir, s) for _k, _a, s in imgs}), images_dir))
    ver = plugin_version()
    tp = set(re.findall(r"Technology Preview v?(\d+\.\d+)", "\n".join(lines)))
    if ver and tp - {ver}:
        print("warning: the guide says Technology Preview %s but CONFIG.VERSION is %s"
              % (", ".join(sorted(tp)), ver))
    return 0


if __name__ == "__main__":
    sys.exit(main())
